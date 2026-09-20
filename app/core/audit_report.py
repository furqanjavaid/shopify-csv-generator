"""
CRO Audit Word report — score /10, conversion-priority findings.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from urllib.parse import urlparse

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.core.audit_templates import (
    CHECK_ID_TO_PASSING,
    CHECK_ID_TO_TEMPLATE,
    EXEC_SUMMARY,
    ISSUE_TEMPLATES,
    PASSING_MESSAGES,
)

FONT = "Arial"
RGB_DARK = RGBColor(0x1A, 0x1A, 0x2E)
RGB_MID = RGBColor(0x2C, 0x3E, 0x50)
RGB_GREY = RGBColor(0x7F, 0x8C, 0x8D)
RGB_RED = RGBColor(0xC0, 0x39, 0x2B)
RGB_ORANGE = RGBColor(0xE6, 0x7E, 0x22)
RGB_GREEN = RGBColor(0x27, 0xAE, 0x60)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CAT_LABELS = {
    "product_page": "Product Page",
    "cart_checkout": "Cart & Checkout",
    "mobile": "Mobile",
    "speed": "Speed",
    "marketing": "Marketing",
}


def _get_issue_text(issue_key: str, evidence: str, url: str) -> dict:
    """Get expert template for an issue, injecting real evidence."""
    template = ISSUE_TEMPLATES.get(issue_key, {})
    return {
        "title": template.get("title", issue_key),
        "severity": template.get("severity", "MEDIUM"),
        "category": template.get("category", "General"),
        "why": template.get("why", "").strip(),
        "fix": template.get("fix", "").strip(),
        "impact": template.get("impact", ""),
        "evidence": evidence,
        "url": url,
    }


def resolve_template_key(finding: dict, raw: dict | None = None) -> str | None:
    """Map a report finding to an ISSUE_TEMPLATES key."""
    cid = finding.get("check_id")
    raw = raw or {}

    if cid == 202:
        policies = raw.get("policy_pages") or {}
        shipping = policies.get("/policies/shipping-policy") or {}
        if not shipping.get("exists"):
            return "no_shipping_policy"
        issue = (finding.get("issue") or "").lower()
        if "thin" in issue or any(
            (info or {}).get("thin") for info in policies.values()
        ):
            return "thin_policies"
        if not finding.get("passed"):
            return "no_shipping_policy"
        return None

    if cid == 7 and finding.get("passed"):
        return None

    return CHECK_ID_TO_TEMPLATE.get(cid)


def enrich_finding(finding: dict, raw: dict | None = None) -> dict:
    """Overlay expert template title/why/fix/impact onto a failing finding."""
    if finding.get("passed") and finding.get("status") == "pass":
        return finding

    key = resolve_template_key(finding, raw)
    if not key or key not in ISSUE_TEMPLATES:
        return finding

    urls = finding.get("urls") or []
    url = urls[0] if urls else ""
    expert = _get_issue_text(key, finding.get("evidence", ""), url)
    enriched = dict(finding)
    enriched["issue"] = expert["title"]
    enriched["severity"] = expert["severity"] or finding.get("severity")
    enriched["impact"] = expert["impact"] or finding.get("impact", "")
    enriched["fix"] = expert["fix"] or finding.get("fix", "")
    enriched["why"] = expert["why"]
    enriched["template_key"] = key
    enriched["check_name"] = expert["title"]
    return enriched


def format_passing_message(finding: dict, raw: dict | None = None) -> str:
    """Build expert passing-check copy when a template exists."""
    raw = raw or {}
    cid = finding.get("check_id")
    key = CHECK_ID_TO_PASSING.get(cid)
    if not key or key not in PASSING_MESSAGES:
        return f"✓ {finding.get('check_name')}: {(finding.get('evidence') or '')[:80]}"

    tpl = PASSING_MESSAGES[key]
    try:
        msg = tpl.format(
            count=raw.get("trust_badge_count", raw.get("email_capture_count", 0)),
            match=raw.get("trust_text_match") or "detected signals",
            price=raw.get("price_text") or "price",
            time=raw.get("homepage_load_time", "—"),
            title=raw.get("meta_title") or "",
        )
    except Exception:
        msg = tpl
    return f"✓ {msg}"


def set_cell_bg(cell, hex_color: str):
    hex_color = hex_color.lstrip("#")
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_run(para, text, size=10, bold=False, color=None, italic=False):
    r = para.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return r


def add_para(doc, text="", size=10, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        add_run(p, text, size=size, bold=bold, color=color)
    return p


def add_section_bar(doc, text, bg="1A1A2E"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg)
    p = cell.paragraphs[0]
    r = p.add_run(f"  {text}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = FONT
    r.font.color.rgb = RGB_WHITE
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_image_safe(doc, path, width=5.0):
    if path and os.path.exists(path):
        try:
            doc.add_picture(path, width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
        except Exception:
            pass
    return False


def score_color(score: float):
    if score >= 7.5:
        return RGB_GREEN, "27AE60"
    if score >= 5:
        return RGB_ORANGE, "E67E22"
    return RGB_RED, "C0392B"


def category_status(score: float) -> str:
    if score >= 7.5:
        return "✓ Good"
    if score >= 5:
        return "⚠ Fix"
    return "✗ Critical"


def make_category_chart(scores: dict, output_dir: str) -> str:
    cats = scores.get("categories") or {}
    labels = [CAT_LABELS.get(k, k) for k in cats]
    values = [cats[k] for k in cats]
    colors = ["#27AE60" if v >= 7.5 else "#E67E22" if v >= 5 else "#C0392B" for v in values]

    fig, ax = plt.subplots(figsize=(6.5, max(2.8, len(labels) * 0.55)))
    fig.patch.set_facecolor("white")
    ax.barh(labels, values, color=colors, height=0.55)
    ax.set_xlim(0, 10)
    ax.set_xlabel("Score / 10")
    ax.set_title("CRO Category Scores")
    for i, v in enumerate(values):
        ax.text(min(v + 0.15, 9.5), i, f"{v}", va="center", fontsize=9)
    path = os.path.join(output_dir, "chart_cro_scores.png")
    fig.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)
    return path


def make_score_badge(score: float, output_dir: str) -> str:
    _, hex_c = score_color(score)
    fig, ax = plt.subplots(figsize=(2.4, 2.4))
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    circle = plt.Circle((0.5, 0.5), 0.42, color=f"#{hex_c}")
    ax.add_patch(circle)
    ax.text(0.5, 0.55, f"{score}", ha="center", va="center", fontsize=28, fontweight="bold", color="white")
    ax.text(0.5, 0.32, "/ 10", ha="center", va="center", fontsize=11, color="white")
    path = os.path.join(output_dir, "chart_health_score.png")
    fig.savefig(path, dpi=140, bbox_inches="tight")
    plt.close(fig)
    return path


def generate_report(audit_data: dict, output_dir: str) -> str:
    store_url = audit_data.get("store_url", "")
    audit_date = audit_data.get("audit_date", datetime.now().isoformat())
    raw = audit_data.get("raw_findings") or {}
    findings_raw = audit_data.get("findings") or []
    findings = [enrich_finding(f, raw) for f in findings_raw]

    # Inject carousel as a medium issue when present and not already covered
    if raw.get("has_carousel") and not any(
        f.get("template_key") == "carousel_risk" for f in findings
    ):
        expert = _get_issue_text(
            "carousel_risk",
            f"Carousel/slider detected · slides={raw.get('carousel_slide_count', 0)}",
            store_url,
        )
        findings.append({
            "check_id": 15,
            "check_name": expert["title"],
            "severity": expert["severity"],
            "category": "speed",
            "passed": False,
            "status": "warning",
            "issue": expert["title"],
            "evidence": expert["evidence"],
            "impact": expert["impact"],
            "fix": expert["fix"],
            "why": expert["why"],
            "urls": [store_url],
            "screenshot": (raw.get("screenshots") or {}).get("homepage") or "",
            "template_key": "carousel_risk",
        })

    scores = audit_data.get("scores") or {"categories": {}, "overall": 0}
    screenshot_dir = os.path.join(output_dir, "screenshots")

    domain = urlparse(store_url).netloc.replace("www.", "")
    store_name = domain.split(".")[0].title() if domain else "Store"
    try:
        date_str = datetime.fromisoformat(audit_date).strftime("%B %d, %Y")
        file_date = datetime.fromisoformat(audit_date).strftime("%Y-%m-%d")
    except Exception:
        date_str = audit_date
        file_date = datetime.now().strftime("%Y-%m-%d")

    overall = float(scores.get("overall") or 0)
    badge = make_score_badge(overall, output_dir)
    chart = make_category_chart(scores, output_dir)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ── PAGE 1 Cover ──
    add_para(doc, "CRO AUDIT REPORT", size=26, bold=True, color=RGB_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, store_name, size=18, bold=True, color=RGB_MID, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, store_url, size=10, color=RGB_GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, date_str, size=10, color=RGB_GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    sc, _ = score_color(overall)
    add_para(doc, f"Overall CRO Score: {overall}/10", size=20, bold=True, color=sc, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_image_safe(doc, badge, width=2.0)
    add_para(doc, "Prepared by: Sentivo Limited", size=10, color=RGB_GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    home_ss = os.path.join(screenshot_dir, "homepage_desktop.png")
    add_image_safe(doc, home_ss, width=5.5)
    doc.add_page_break()

    # ── PAGE 2 Executive Summary ──
    add_section_bar(doc, "EXECUTIVE SUMMARY")
    add_para(doc, f"Overall CRO Score: {overall}/10", size=14, bold=True, color=sc)

    highs = [f for f in findings if f.get("severity") == "HIGH" and not f.get("passed")]
    mediums = [f for f in findings if f.get("severity") == "MEDIUM" and not f.get("passed")]
    lows = [f for f in findings if f.get("severity") == "LOW" and not f.get("passed")]

    add_para(doc, "Top critical issues", size=12, bold=True, color=RGB_RED, space_after=4)
    if highs:
        for f in highs[:3]:
            add_para(doc, f"• {f.get('issue', '')}", size=10, space_after=2)
    else:
        add_para(doc, "• No HIGH severity issues found.", size=10, color=RGB_GREEN)

    add_para(doc, "Quick wins", size=12, bold=True, color=RGB_ORANGE, space_after=4)
    if mediums:
        for f in mediums[:4]:
            add_para(doc, f"• {f.get('issue', '')}", size=10, space_after=2)
    else:
        add_para(doc, "• No medium-priority gaps flagged.", size=10)

    # Expert executive summary
    if overall >= 7.5:
        summary = EXEC_SUMMARY["score_high"].format(store=store_name).strip()
    elif overall >= 5:
        summary = EXEC_SUMMARY["score_medium"].format(store=store_name).strip()
    else:
        summary = EXEC_SUMMARY["score_low"].format(store=store_name).strip()
    summary = re.sub(r"\s*\n\s*", " ", summary)
    add_para(doc, summary, size=10, space_after=12)
    doc.add_page_break()

    # ── PAGE 3 Score Breakdown ──
    add_section_bar(doc, "SCORE BREAKDOWN")
    add_image_safe(doc, chart, width=5.8)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, label in enumerate(("Category", "Score", "Status")):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_run(p, label, size=10, bold=True, color=RGB_WHITE)
        set_cell_bg(hdr[i], "1A1A2E")

    for key, val in (scores.get("categories") or {}).items():
        row = table.add_row().cells
        row[0].text = CAT_LABELS.get(key, key)
        row[1].text = f"{val}/10"
        row[2].text = category_status(val)

    doc.add_paragraph()
    doc.add_page_break()

    # ── Detailed findings ──
    add_section_bar(doc, "DETAILED FINDINGS", bg="C0392B")
    add_para(doc, "🔴 CRITICAL ISSUES (HIGH)", size=14, bold=True, color=RGB_RED)

    def write_full_issue(f: dict):
        add_para(doc, f.get("issue", ""), size=11, bold=True, space_after=2)
        add_para(doc, f"Evidence: {f.get('evidence', '')}", size=9, color=RGB_GREY, space_after=1)
        urls = f.get("urls") or []
        if urls:
            add_para(doc, f"URL(s): {', '.join(urls[:4])}", size=8, color=RGB_GREY, space_after=1)
        if f.get("why"):
            add_para(doc, "Why this matters", size=10, bold=True, space_after=1)
            add_para(doc, f["why"], size=9, space_after=4)
        elif f.get("impact"):
            add_para(doc, f"Why it matters: {f['impact']}", size=9, space_after=1)
        if f.get("fix"):
            add_para(doc, "How to fix", size=10, bold=True, space_after=1)
            add_para(doc, f["fix"], size=9, space_after=2)
        if f.get("impact") and f.get("why"):
            add_para(doc, f"Expected impact: {f['impact']}", size=9, bold=True, color=RGB_MID, space_after=4)
        shot = f.get("screenshot") or ""
        if shot:
            add_image_safe(doc, os.path.join(screenshot_dir, shot), width=4.8)

    if highs:
        for f in highs:
            write_full_issue(f)
    else:
        add_para(doc, "No critical issues.", size=10, color=RGB_GREEN)

    doc.add_page_break()
    add_section_bar(doc, "IMPROVEMENTS", bg="E67E22")
    add_para(doc, "🟡 IMPROVEMENTS (MEDIUM)", size=14, bold=True, color=RGB_ORANGE)
    if mediums:
        for f in mediums:
            write_full_issue(f)
    else:
        add_para(doc, "No medium issues.", size=10, color=RGB_GREEN)

    doc.add_page_break()
    add_section_bar(doc, "NICE TO HAVE", bg="27AE60")
    add_para(doc, "🟢 NICE TO HAVE (LOW)", size=14, bold=True, color=RGB_GREEN)
    if lows:
        for f in lows:
            write_full_issue(f)
    else:
        add_para(doc, "• No low-priority gaps.", size=10)

    # Passed checks brief — expert passing messages
    passes = [f for f in findings if f.get("passed") and f.get("status") == "pass"]
    if passes:
        add_para(doc, "Passing checks", size=12, bold=True, space_after=4)
        for f in passes:
            add_para(doc, format_passing_message(f, raw), size=8, color=RGB_GREY, space_after=1)

    # ── Multi-page site health sections ──
    doc.add_page_break()
    add_section_bar(doc, "SITE HEALTH — MULTI-PAGE CRAWL", bg="2C3E50")

    pages = audit_data.get("pages_crawled") or []
    add_para(doc, f"Pages crawled: {len(pages)}", size=10, bold=True, space_after=4)
    for u in pages[:12]:
        add_para(doc, f"• {u}", size=8, color=RGB_GREY, space_after=1)

    # Dead Links
    add_para(doc, "Dead Links", size=12, bold=True, color=RGB_RED, space_after=4)
    dead = audit_data.get("dead_links") or raw.get("dead_links") or []
    if dead:
        for link in dead[:15]:
            add_para(doc, f"• 404: {link}", size=9, space_after=2)
    else:
        add_para(doc, f"• {PASSING_MESSAGES['dead_links']}", size=9, color=RGB_GREEN)

    # Policy Pages
    add_para(doc, "Policy Pages", size=12, bold=True, color=RGB_ORANGE, space_after=4)
    policies = audit_data.get("policy_pages") or raw.get("policy_pages") or {}
    if policies:
        for path, info in policies.items():
            if not info.get("exists"):
                add_para(doc, f"• {path} — missing", size=9, color=RGB_RED, space_after=2)
            elif info.get("thin"):
                add_para(
                    doc,
                    f"• {path} — thin content ({info.get('word_count', 0)} words)",
                    size=9,
                    color=RGB_ORANGE,
                    space_after=2,
                )
            else:
                add_para(
                    doc,
                    f"• {path} — OK ({info.get('word_count', 0)} words)",
                    size=9,
                    color=RGB_GREEN,
                    space_after=2,
                )
    else:
        add_para(doc, "• Policy scan unavailable.", size=9, color=RGB_GREY)

    # Navigation
    add_para(doc, "Navigation", size=12, bold=True, space_after=4)
    nav = audit_data.get("nav_summary") or {}
    nav_count = nav.get("nav_link_count", raw.get("nav_links_count", 0))
    hamburger = nav.get("nav_hamburger_only", raw.get("nav_hamburger_only", False))
    if hamburger:
        expert = ISSUE_TEMPLATES["hamburger_only_desktop"]
        add_para(doc, f"• {expert['title']}", size=9, color=RGB_ORANGE, space_after=2)
        add_para(doc, f"  Visible links: {nav_count}", size=8, color=RGB_GREY, space_after=2)
    else:
        add_para(
            doc,
            f"• {PASSING_MESSAGES['nav_depth'].format(count=nav_count)}",
            size=9,
            color=RGB_GREEN,
            space_after=2,
        )

    # Price Consistency
    add_para(doc, "Price Consistency", size=12, bold=True, space_after=4)
    mismatches = audit_data.get("price_mismatches") or raw.get("price_mismatches") or []
    if mismatches:
        expert = ISSUE_TEMPLATES["price_mismatch"]
        add_para(doc, f"• {expert['title']}", size=9, color=RGB_RED, space_after=2)
        for m in mismatches[:8]:
            add_para(
                doc,
                f"  {m.get('product_url', '')}: collection={m.get('collection_price')} "
                f"vs PDP={m.get('pdp_price')}",
                size=8,
                color=RGB_GREY,
                space_after=2,
            )
    else:
        add_para(
            doc,
            "• No collection vs PDP price mismatches in sampled products.",
            size=9,
            color=RGB_GREEN,
        )

    doc.add_page_break()
    add_section_bar(doc, "ABOUT SENTIVO")
    add_para(
        doc,
        "This audit was prepared by Sentivo Limited.\n"
        "For implementation support, contact: hello@sentivo.co",
        size=11,
        space_after=8,
    )
    add_para(doc, store_url, size=9, color=RGB_GREY)

    safe_domain = re.sub(r"[^\w\-]", "_", domain)
    out_name = f"CRO_Audit_{safe_domain}_{file_date}.docx"
    out_path = os.path.join(output_dir, out_name)
    doc.save(out_path)
    return out_path


# Back-compat aliases used by older code paths
def deduplicate_results(results: list) -> dict:
    return {}


def compute_cro_scores(deduped: dict) -> dict:
    return {"categories": {}, "overall": 0}
