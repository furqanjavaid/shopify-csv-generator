"""Site health orchestration and reporting (audit → pilot → full → validation)."""

from __future__ import annotations

import json
import logging
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sentivo_extractor.core.input_csv import read_seed_csv
from sentivo_extractor.core.production_validation import (
    STRICT_SUCCESS_RATE_MIN,
    build_validation_summary,
    strict_validation_passed,
)
from sentivo_extractor.core.site_rule_suggester import domain_from_url
from sentivo_extractor.core.utils import write_json

logger = logging.getLogger("sentivo_extractor")

HEALTH_COLUMNS = [
    "Domain",
    "Platform",
    "Products Expected",
    "Products Extracted",
    "Success Rate",
    "Validation Passed",
    "Images %",
    "Variants %",
    "Price %",
    "SKU %",
    "Retry Count",
    "Failed Products",
    "Needs YAML",
    "Recommended Action",
]

ACTION_READY = "Ready for Production"
ACTION_YAML = "Improve YAML"
ACTION_DISCOVERY = "Improve Product Discovery"
ACTION_PDP = "Improve PDP Extraction"
ACTION_VARIANT = "Improve Variant Engine"
ACTION_IMAGE = "Improve Image Engine"
ACTION_MANUAL = "Manual Investigation"


def run_site_health_pipeline(
    input_csv: Path,
    output_root: Path,
    options: dict[str, Any],
) -> dict[str, Any]:
    """
    Run audit, pilot extract, full extract, and production-validation extract.
    Does not modify extraction logic — only invokes existing auditors/crawlers.
    """
    from sentivo_extractor.core.audit import DomainAuditor
    from sentivo_extractor.core.crawler import UniversalCrawler

    output_root = Path(output_root)
    output_root.mkdir(parents=True, exist_ok=True)
    input_csv = Path(input_csv)
    logger.info("Product Discovery Version: Unified")

    phases: dict[str, str] = {}

    audit_opts = dict(options)
    audit_opts["output"] = str(output_root / "audit")
    phases["audit"] = audit_opts["output"]
    logger.info("Site health phase 1/4: audit")
    DomainAuditor(audit_opts).run(input_csv)

    pilot_opts = dict(options)
    pilot_opts["output"] = str(output_root / "pilot")
    pilot_opts["pilot"] = True
    pilot_opts["production_validation"] = False
    pilot_opts["strict"] = False
    pilot_opts["overwrite"] = True
    phases["pilot"] = pilot_opts["output"]
    logger.info("Site health phase 2/4: pilot")
    UniversalCrawler(pilot_opts).run(input_csv)

    full_opts = dict(options)
    full_opts["output"] = str(output_root / "full")
    full_opts["pilot"] = False
    full_opts["production_validation"] = False
    full_opts["strict"] = False
    full_opts["overwrite"] = True
    phases["full"] = full_opts["output"]
    logger.info("Site health phase 3/4: full extraction")
    UniversalCrawler(full_opts).run(input_csv)

    val_opts = dict(options)
    val_opts["output"] = str(output_root / "validation")
    val_opts["pilot"] = False
    val_opts["production_validation"] = True
    val_opts["strict"] = False
    val_opts["overwrite"] = True
    phases["validation"] = val_opts["output"]
    logger.info("Site health phase 4/4: production validation")
    UniversalCrawler(val_opts).run(input_csv)

    report = build_site_health_report(
        output_root=output_root,
        input_csv=input_csv,
        phases=phases,
    )
    write_site_health_outputs(output_root, report)
    return report


def build_site_health_report(
    *,
    output_root: Path,
    input_csv: Path | None = None,
    phases: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Aggregate domain health from phase output directories."""
    output_root = Path(output_root)
    phases = phases or {
        "audit": str(output_root / "audit"),
        "pilot": str(output_root / "pilot"),
        "full": str(output_root / "full"),
        "validation": str(output_root / "validation"),
    }

    expected_by_domain: dict[str, Any] = {}
    if input_csv and Path(input_csv).exists():
        for seed in read_seed_csv(Path(input_csv)):
            dom = domain_from_url(seed["url"])
            raw = (seed.get("expected_count") or "").strip()
            if raw:
                try:
                    expected_by_domain[dom] = int(float(raw))
                except ValueError:
                    expected_by_domain[dom] = raw

    audit_rows = _load_audit_domain_rows(Path(phases["audit"]))
    audit_by_dom = {str(r.get("domain") or ""): r for r in audit_rows}

    full_summary = _load_json(Path(phases["full"]) / "run_summary.json")
    full_cov = {
        str(r.get("domain") or ""): r for r in (full_summary.get("coverage") or [])
    }

    validation_products = _load_validation_products(Path(phases["validation"]))
    retry_by_dom = _retry_counts_by_domain(Path(phases["validation"]) / "retry_queue.json")

    domains: set[str] = set()
    domains.update(audit_by_dom.keys())
    domains.update(full_cov.keys())
    domains.update(validation_products.keys())
    domains.update(expected_by_domain.keys())
    domains.discard("")

    domain_rows: list[dict[str, Any]] = []
    for dom in sorted(domains):
        audit = audit_by_dom.get(dom) or {}
        cov = full_cov.get(dom) or {}
        val_rows = validation_products.get(dom) or []
        val_summary = build_validation_summary(val_rows) if val_rows else {}

        products_expected = (
            expected_by_domain.get(dom)
            or audit.get("expected_count")
            or cov.get("expected_count")
            or ""
        )
        products_extracted = cov.get("extracted_count") or audit.get("extracted_count") or 0
        try:
            products_extracted = int(products_extracted)
        except (TypeError, ValueError):
            products_extracted = 0

        success_rate = float(val_summary.get("Success Rate %") or 0)
        if not val_rows and products_extracted:
            success_rate = 100.0

        field_rates = _field_pass_rates(val_rows)
        retry_count = retry_by_dom.get(dom, 0)
        if val_rows:
            retry_count = sum(int(r.get("Retry Count") or 0) for r in val_rows)

        failed_products = int(val_summary.get("Failed") or 0)
        needs_yaml = _needs_yaml(audit, cov, val_rows)

        row = {
            "Domain": dom,
            "Platform": audit.get("platform_detected") or "Unknown",
            "Products Expected": products_expected,
            "Products Extracted": products_extracted,
            "Success Rate": success_rate,
            "Validation Passed": "yes"
            if strict_validation_passed(val_summary)
            else "no",
            "Images %": field_rates["images"],
            "Variants %": field_rates["variants"],
            "Price %": field_rates["price"],
            "SKU %": field_rates["sku"],
            "Retry Count": retry_count,
            "Failed Products": failed_products,
            "Needs YAML": needs_yaml,
            "Recommended Action": recommend_action(
                platform=str(audit.get("platform_detected") or ""),
                success_rate=success_rate,
                validation_passed=strict_validation_passed(val_summary),
                failed_products=failed_products,
                needs_yaml=needs_yaml,
                products_extracted=products_extracted,
                discovered=int(cov.get("discovered_count") or audit.get("product_urls_found") or 0),
                field_rates=field_rates,
                audit=audit,
                val_rows=val_rows,
            ),
        }
        domain_rows.append(row)

    summary = {
        "domains_total": len(domain_rows),
        "validation_passed_domains": sum(
            1 for r in domain_rows if r.get("Validation Passed") == "yes"
        ),
        "ready_for_production": sum(
            1 for r in domain_rows if r.get("Recommended Action") == ACTION_READY
        ),
        "needs_yaml_domains": sum(
            1 for r in domain_rows if str(r.get("Needs YAML")).lower() in {"yes", "true"}
        ),
        "average_success_rate": round(
            sum(float(r.get("Success Rate") or 0) for r in domain_rows) / len(domain_rows),
            2,
        )
        if domain_rows
        else 0.0,
        "strict_threshold_percent": STRICT_SUCCESS_RATE_MIN,
    }

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_csv": str(input_csv or ""),
        "phases": phases,
        "domains": domain_rows,
        "summary": summary,
    }


def recommend_action(
    *,
    platform: str,
    success_rate: float,
    validation_passed: bool,
    failed_products: int,
    needs_yaml: str,
    products_extracted: int,
    discovered: int,
    field_rates: dict[str, float],
    audit: dict[str, Any],
    val_rows: list[dict[str, Any]],
) -> str:
    """Heuristic weakness identification — no extraction changes."""
    if str(needs_yaml).lower() in {"yes", "true"}:
        if discovered <= 0 or products_extracted <= 0:
            return ACTION_YAML
        if float(audit.get("title_success_rate") or 0) < 0.5:
            return ACTION_YAML

    if validation_passed and failed_products == 0 and success_rate >= STRICT_SUCCESS_RATE_MIN:
        if str(needs_yaml).lower() not in {"yes", "true"}:
            return ACTION_READY

    if discovered <= 0 or (products_extracted <= 0 and discovered <= 5):
        return ACTION_DISCOVERY

    missing_counts = _missing_field_counts(val_rows)
    if field_rates.get("images", 100) < 90 or missing_counts.get("images", 0) >= max(
        1, len(val_rows) // 4
    ):
        return ACTION_IMAGE
    if field_rates.get("variants", 100) < 90 or missing_counts.get("variants", 0) >= max(
        1, len(val_rows) // 4
    ):
        return ACTION_VARIANT

    if (
        field_rates.get("price", 100) < 90
        or field_rates.get("sku", 100) < 90
        or missing_counts.get("price", 0)
        or missing_counts.get("title", 0)
    ):
        return ACTION_PDP

    if success_rate < STRICT_SUCCESS_RATE_MIN or failed_products > 0:
        return ACTION_MANUAL

    return ACTION_MANUAL


def write_site_health_outputs(output_root: Path, report: dict[str, Any]) -> tuple[Path, Path]:
    output_root = Path(output_root)
    xlsx_path = output_root / "site_health_report.xlsx"
    json_path = output_root / "site_health_summary.json"

    rows = report.get("domains") or []
    _write_health_xlsx(xlsx_path, rows, report.get("summary") or {})
    write_json(json_path, report)
    return xlsx_path, json_path


def _write_health_xlsx(
    path: Path, rows: list[dict[str, Any]], summary: dict[str, Any]
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        import openpyxl

        wb = openpyxl.Workbook()
        ws_sum = wb.active
        ws_sum.title = "Summary"
        ws_sum.append(["metric", "value"])
        for key, val in summary.items():
            ws_sum.append([key, val])

        ws = wb.create_sheet("Site Health")
        ws.append(HEALTH_COLUMNS)
        for row in rows:
            ws.append([row.get(c, "") for c in HEALTH_COLUMNS])
        wb.save(path)
        wb.close()
        return path
    except Exception:
        import csv

        csv_path = path.with_suffix(".csv")
        with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=HEALTH_COLUMNS, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow(row)
        return csv_path


def _load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception:
        return {}


def _load_audit_domain_rows(audit_dir: Path) -> list[dict[str, Any]]:
    xlsx = audit_dir / "domain_audit.xlsx"
    csv_path = audit_dir / "domain_audit.csv"
    if xlsx.exists():
        try:
            import openpyxl

            wb = openpyxl.load_workbook(xlsx, read_only=True)
            ws = wb.active
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            rows: list[dict[str, Any]] = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({str(headers[i]): row[i] for i in range(len(headers))})
            wb.close()
            return rows
        except Exception:
            pass
    if csv_path.exists():
        import csv

        with csv_path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    return []


def _load_validation_products(validation_dir: Path) -> dict[str, list[dict[str, Any]]]:
    path = validation_dir / "final_validation_report.xlsx"
    json_path = validation_dir / "final_validation_report.json"
    products: list[dict[str, Any]] = []

    if path.exists():
        try:
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True)
            if "Products" in wb.sheetnames:
                ws = wb["Products"]
            else:
                ws = wb.active
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                products.append(
                    {str(headers[i]): row[i] for i in range(len(headers)) if headers[i]}
                )
            wb.close()
        except Exception:
            pass
    elif json_path.exists():
        data = _load_json(json_path)
        products = list(data.get("products") or [])

    by_dom: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in products:
        url = str(row.get("URL") or "")
        by_dom[domain_from_url(url)].append(row)
    return by_dom


def _retry_counts_by_domain(path: Path) -> dict[str, int]:
    data = _load_json(path)
    items = data.get("retry_queue") or []
    counts: dict[str, int] = defaultdict(int)
    for item in items:
        url = str(item.get("url") or "")
        counts[domain_from_url(url)] += 1
    return dict(counts)


def _field_pass_rates(rows: list[dict[str, Any]]) -> dict[str, float]:
    if not rows:
        return {"images": 0.0, "variants": 0.0, "price": 0.0, "sku": 0.0}
    total = len(rows)
    ok = {"images": 0, "variants": 0, "price": 0, "sku": 0}
    for row in rows:
        reason = str(row.get("Failure Reason") or "").lower()
        if int(row.get("Images") or 0) > 0 and "images" not in reason:
            ok["images"] += 1
        if int(row.get("Variants") or 0) > 0 and "variants" not in reason:
            ok["variants"] += 1
        if str(row.get("Price") or "").strip() and "price" not in reason:
            ok["price"] += 1
        if str(row.get("SKU") or "").strip() and "sku" not in reason:
            ok["sku"] += 1
    return {k: round(ok[k] / total * 100.0, 1) for k in ok}


def _missing_field_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        reason = str(row.get("Failure Reason") or "").lower()
        for field in ("title", "price", "images", "variants", "description", "sku"):
            if field in reason:
                counts[field] += 1
    return dict(counts)


def _needs_yaml(
    audit: dict[str, Any],
    cov: dict[str, Any],
    val_rows: list[dict[str, Any]],
) -> str:
    flag = str(audit.get("needs_yaml_rules") or "").lower()
    if flag in {"yes", "true", "1"}:
        return "yes"
    discovered = int(cov.get("discovered_count") or audit.get("product_urls_found") or 0)
    if discovered <= 0:
        return "yes"
    if val_rows and all(str(r.get("Status") or "") == "Failed" for r in val_rows):
        notes = str(audit.get("notes") or "").lower()
        if "yaml" in notes or "selector" in notes:
            return "yes"
    return "no"
